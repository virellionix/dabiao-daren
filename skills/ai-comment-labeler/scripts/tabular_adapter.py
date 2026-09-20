#!/usr/bin/env python3
"""Generic CSV/XLSX boundary for the labeling Skill.

The adapter is intentionally thin: it maps caller-selected columns to the
shared JSONL contract and maps a complete result back to the original rows.
It does not contain a product's label tree, infer IDs, or call a model.
CSV/TSV use the standard library. XLSX support is optional and uses
``openpyxl`` only when the calling environment has chosen to install it.
"""

from __future__ import annotations

import argparse
import csv
import json
import sys
from pathlib import Path

try:
    from shared_labeler import adapt_record
except ImportError:  # pragma: no cover - supports direct module loading
    sys.path.insert(0, str(Path(__file__).resolve().parent))
    from shared_labeler import adapt_record


BASE_COLUMNS = ("打标状态", "复核原因", "AI标签JSON", "证据JSON", "行为JSON")


def _clean(value):
    return "" if value is None else str(value)


def parse_mapping(value: str) -> tuple[str, str]:
    if "=" not in value:
        raise ValueError("mapping must be name=column: " + value)
    name, column = value.split("=", 1)
    if not name or not column:
        raise ValueError("mapping must have non-empty name and column: " + value)
    return name, column


def read_rows(path: str | Path) -> tuple[list[str], list[dict]]:
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix in {".xlsx", ".xlsm"}:
        try:
            import openpyxl
        except ImportError as exc:
            raise RuntimeError("读取 XLSX 需要在调用方环境安装 openpyxl；CSV 不需要第三方依赖") from exc
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        sheet = workbook.active
        iterator = sheet.iter_rows(values_only=True)
        try:
            header_values = next(iterator)
        except StopIteration:
            return [], []
        headers = [_clean(value).strip() for value in header_values]
        rows = [{headers[index]: _clean(value) for index, value in enumerate(values)
                 if index < len(headers) and headers[index]}
                for values in iterator]
        return _validate_headers(headers), rows
    delimiter = "\t" if suffix == ".tsv" else ","
    with path.open(newline="", encoding="utf-8-sig") as stream:
        reader = csv.DictReader(stream, delimiter=delimiter)
        headers = _validate_headers(reader.fieldnames or [])
        rows = [{key: _clean(value) for key, value in row.items() if key is not None} for row in reader]
    return headers, rows


def _validate_headers(headers) -> list[str]:
    result = [_clean(item).strip() for item in headers]
    if not result or any(not item for item in result):
        raise ValueError("tabular input needs a non-empty header row")
    if len(result) != len(set(result)):
        raise ValueError("tabular input has duplicate column names")
    return result


def write_rows(path: str | Path, headers: list[str], rows: list[dict]) -> None:
    path = Path(path)
    suffix = path.suffix.lower()
    headers = _validate_headers(headers)
    if suffix in {".xlsx", ".xlsm"}:
        try:
            import openpyxl
        except ImportError as exc:
            raise RuntimeError("写入 XLSX 需要在调用方环境安装 openpyxl；CSV 不需要第三方依赖") from exc
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.append(headers)
        for row in rows:
            sheet.append([row.get(header, "") for header in headers])
        workbook.save(path)
        return
    delimiter = "\t" if suffix == ".tsv" else ","
    with path.open("w", newline="", encoding="utf-8-sig") as stream:
        writer = csv.DictWriter(stream, fieldnames=headers, delimiter=delimiter, extrasaction="ignore")
        writer.writeheader()
        writer.writerows({header: _clean(row.get(header, "")) for header in headers} for row in rows)


def prepare_rows(rows: list[dict], *, id_column: str, text_column: str,
                 context_columns=None, source_kind="full_text") -> list[dict]:
    context_columns = dict(context_columns or {})
    return [adapt_record(row, id_column=id_column, text_column=text_column,
                         context_columns=context_columns, source_kind=source_kind)
            for row in rows]


def write_jsonl(path: str | Path, rows: list[dict]) -> None:
    Path(path).write_text("".join(json.dumps(row, ensure_ascii=False, sort_keys=True) + "\n"
                                         for row in rows), encoding="utf-8")


def read_jsonl(path: str | Path) -> list[dict]:
    rows = []
    with Path(path).open(encoding="utf-8") as stream:
        for line in stream:
            if line.strip():
                rows.append(json.loads(line))
    return rows


def _result_status(result: dict) -> str:
    if result.get("status") in {"candidate", "review", "error"}:
        return result["status"]
    if result.get("review_required") is True:
        return "review"
    if result.get("prediction") is not None:
        return "candidate"
    return "error"


def _result_reason(result: dict) -> str:
    details = result.get("review_details")
    if isinstance(details, list):
        return "；".join(str(item.get("detail", item)) if isinstance(item, dict) else str(item)
                         for item in details)
    reasons = result.get("review_reasons")
    if isinstance(reasons, list):
        return "；".join(str(item) for item in reasons)
    return ""


def _labels(prediction: dict | None, project: dict | None) -> dict:
    if not isinstance(prediction, dict):
        return {}
    labels = prediction.get("labels")
    if not isinstance(labels, dict):
        return {"stance": prediction.get("stance")} if prediction.get("stance") else {}
    output = {}
    dimensions = (project or {}).get("dimensions", {})
    for dimension, values in labels.items():
        names = []
        for item in values if isinstance(values, list) else []:
            value = item.get("value") if isinstance(item, dict) else item
            names.append(dimensions.get(dimension, {}).get("labels", {}).get(value, {}).get("name", value))
        output[dimension] = names
    return output


def merge_rows(original: list[dict], results: list[dict], *, id_column="id", project=None,
               dimension_columns=None) -> tuple[list[str], list[dict]]:
    if id_column not in (list(original[0]) if original else []):
        raise ValueError("id column not found in original input: " + id_column)
    by_id = {str(item.get("id")): item for item in results}
    if len(by_id) != len(results):
        raise ValueError("result JSONL contains duplicate IDs")
    id_to_row = {str(item.get(id_column)): item for item in original}
    missing = set(id_to_row) - set(by_id)
    extra = set(by_id) - set(id_to_row)
    if missing or extra:
        raise ValueError("result/input ID mismatch: missing=%s extra=%s" %
                         (len(missing), len(extra)))
    dimension_columns = dict(dimension_columns or {})
    headers = list(original[0]) if original else []
    for column in BASE_COLUMNS + tuple(dimension_columns.values()):
        if column not in headers:
            headers.append(column)
    output = []
    for source in original:
        row = dict(source)
        result = by_id[str(source[id_column])]
        prediction = result.get("prediction")
        labels = _labels(prediction, project)
        row["打标状态"] = _result_status(result)
        row["复核原因"] = _result_reason(result)
        row["AI标签JSON"] = json.dumps(labels, ensure_ascii=False, sort_keys=True)
        row["证据JSON"] = json.dumps((prediction or {}).get("evidence", []), ensure_ascii=False)
        row["行为JSON"] = json.dumps((prediction or {}).get("behaviors", []), ensure_ascii=False)
        for dimension, column in dimension_columns.items():
            row[column] = "、".join(labels.get(dimension, []))
        output.append(row)
    return headers, output


def _prepare_cli(args) -> int:
    headers, rows = read_rows(args.input)
    if args.id_column not in headers or args.text_column not in headers:
        raise ValueError("id/text column not found in input header")
    context = dict(parse_mapping(value) for value in args.context_column)
    missing = [column for column in context.values() if column not in headers]
    if missing:
        raise ValueError("context column not found: " + ",".join(missing))
    write_jsonl(args.output, prepare_rows(rows, id_column=args.id_column,
                                          text_column=args.text_column,
                                          context_columns=context,
                                          source_kind=args.source_kind))
    print(json.dumps({"rows": len(rows), "output": str(args.output)}, ensure_ascii=False))
    return 0


def _merge_cli(args) -> int:
    headers, rows = read_rows(args.input)
    project = json.loads(Path(args.project).read_text(encoding="utf-8")) if args.project else None
    dimensions = dict(parse_mapping(value) for value in args.dimension_column)
    out_headers, out_rows = merge_rows(rows, read_jsonl(args.results), project=project,
                                       id_column=args.id_column, dimension_columns=dimensions)
    write_rows(args.output, out_headers, out_rows)
    print(json.dumps({"rows": len(out_rows), "output": str(args.output)}, ensure_ascii=False))
    return 0


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    sub = parser.add_subparsers(dest="command", required=True)
    prepare = sub.add_parser("prepare", help="tabular rows -> generic JSONL input")
    prepare.add_argument("--input", required=True)
    prepare.add_argument("--output", required=True)
    prepare.add_argument("--id-column", required=True)
    prepare.add_argument("--text-column", required=True)
    prepare.add_argument("--context-column", action="append", default=[], metavar="NAME=COLUMN")
    prepare.add_argument("--source-kind", choices=("full_text", "excerpt"), default="full_text")
    prepare.set_defaults(func=_prepare_cli)
    merge = sub.add_parser("merge", help="complete result JSONL -> original tabular shape")
    merge.add_argument("--input", required=True)
    merge.add_argument("--results", required=True)
    merge.add_argument("--output", required=True)
    merge.add_argument("--id-column", required=True)
    merge.add_argument("--project")
    merge.add_argument("--dimension-column", action="append", default=[], metavar="DIMENSION=COLUMN")
    merge.set_defaults(func=_merge_cli)
    args = parser.parse_args(argv)
    try:
        return args.func(args)
    except (OSError, ValueError, RuntimeError, json.JSONDecodeError) as exc:
        parser.error(str(exc))


if __name__ == "__main__":
    raise SystemExit(main())
